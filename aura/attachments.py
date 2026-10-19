"""Files the user attaches to a message: images (read by a vision model) and documents (text extracted).

The main chat model is text-only, so attachments are turned into text once, when the message is sent:
an image's text is transcribed and the rest described; a document's text is extracted. That text goes into
the message, so every tool works with it ("set reminders for my gym schedule") and later turns can refer
to it without re-sending the image. Copies are kept in Data/attachments so history still shows them.
"""
import base64
import io
import re
import shutil
import uuid
from pathlib import Path

from langchain_core.messages import HumanMessage
from langchain_groq import ChatGroq
from pypdf import PdfReader

from . import config

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
DOC_EXT = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".py", ".js", ".ts", ".html", ".css",
           ".rst", ".org", ".tex", ".log", ".yaml", ".yml", ".xml", ".sql", ".c", ".cpp", ".java", ".rs", ".go"}
MAX_BYTES = 25 * 1024 * 1024
MAX_DOC_CHARS = 8000      # per document, so one attachment can't blow the per-minute token budget
MAX_IMAGE_SIDE = 1600     # downscale big photos before sending them to the vision model

VISION_PROMPT = """Transcribe ALL text in this image exactly, keeping its structure (rows of a table or schedule
as separate lines, e.g. "Mon: Chest + Triceps 7:00 PM"). Then add one or two sentences describing anything
else that matters (what kind of image it is, charts, people, objects). Plain text only."""


def folder() -> Path:
    path = config.DATA_DIR / "attachments"
    path.mkdir(parents=True, exist_ok=True)
    return path


def kind_of(name: str) -> str | None:
    ext = Path(name).suffix.lower()
    return "image" if ext in IMAGE_EXT else "document" if ext in DOC_EXT else None


def _meta(path: Path, name: str) -> dict:
    return {"id": path.name, "name": name, "kind": kind_of(name), "path": str(path), "size": path.stat().st_size}


def add_file(source: str) -> dict:
    """Copy a file the user picked into Data/attachments."""
    src = Path(source)
    if not kind_of(src.name):
        raise ValueError(f"{src.name}: only images and documents (PDF, Word, text, code) can be attached.")
    if src.stat().st_size > MAX_BYTES:
        raise ValueError(f"{src.name} is too large (max 25 MB).")
    dest = folder() / f"{uuid.uuid4().hex[:8]}_{_safe(src.name)}"
    shutil.copyfile(src, dest)
    return _meta(dest, src.name)


def add_bytes(name: str, data: bytes) -> dict:
    """A pasted screenshot or a file dropped onto the window."""
    if not kind_of(name):
        raise ValueError(f"{name}: only images and documents can be attached.")
    if len(data) > MAX_BYTES:
        raise ValueError(f"{name} is too large (max 25 MB).")
    dest = folder() / f"{uuid.uuid4().hex[:8]}_{_safe(name)}"
    dest.write_bytes(data)
    return _meta(dest, name)


def get(attachment_id: str) -> dict | None:
    """Look an attachment up by id; only files inside Data/attachments are ever used."""
    path = (folder() / Path(attachment_id).name).resolve()
    if not path.is_file() or not path.is_relative_to(folder().resolve()):
        return None
    return _meta(path, path.name.split("_", 1)[-1])


def _safe(name: str) -> str:
    return re.sub(r"[^\w.\- ]", "_", name)[-80:] or "file"


# ── turning attachments into text ────────────────────────────────────────
def _image_data_url(path: Path) -> str:
    from PIL import Image
    with Image.open(path) as img:
        img = img.convert("RGB")
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=88)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def read_image(path: Path, llm=None) -> str:
    llm = llm or ChatGroq(model=config.VISION_MODEL, api_key=config.GROQ_API_KEY, max_tokens=1200,
                          max_retries=3, request_timeout=90, reasoning_format="hidden")
    message = HumanMessage(content=[{"type": "text", "text": VISION_PROMPT},
                                    {"type": "image_url", "image_url": {"url": _image_data_url(path)}}])
    return str(llm.invoke([message]).content).strip()


def read_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    elif ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for table in doc.tables for row in table.rows]
        text = "\n".join([p.text for p in doc.paragraphs] + rows)
    else:
        text = path.read_bytes().decode("utf-8", errors="replace")
    text = text.strip()
    if len(text) > MAX_DOC_CHARS:
        text = text[:MAX_DOC_CHARS] + f"\n[... truncated: showing {MAX_DOC_CHARS} of {len(text)} characters]"
    return text or "(no readable text: it may be a scanned document)"


def to_text(meta: dict, status=lambda s: None, vision_llm=None) -> str:
    """The text the chat model sees for one attachment."""
    path = Path(meta["path"])
    status(f"reading {meta['name']}")
    try:
        if meta["kind"] == "image":
            body = read_image(path, vision_llm)
            return f"[Attached image: {meta['name']}]\n{body}"
        return f"[Attached document: {meta['name']}]\n{read_document(path)}"
    except Exception as e:
        return f"[Attached {meta['kind']}: {meta['name']} — couldn't be read ({type(e).__name__}). Tell the user.]"

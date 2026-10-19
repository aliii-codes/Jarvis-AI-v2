from pathlib import Path

import docx
import pytest
from langchain_core.messages import AIMessage, HumanMessage

from aura import attachments, config, threads


class FakeVision:
    def __init__(self):
        self.seen = []

    def invoke(self, messages):
        self.seen.append(messages[0].content)
        return AIMessage("GYM SCHEDULE\nMon: Chest 7:00 PM\n\nA typed weekly plan.")


def test_unsupported_and_huge_files_are_refused(tmp_path):
    (tmp_path / "app.exe").write_bytes(b"MZ")
    with pytest.raises(ValueError, match="only images and documents"):
        attachments.add_file(str(tmp_path / "app.exe"))
    with pytest.raises(ValueError, match="too large"):
        attachments.add_bytes("big.png", b"0" * (attachments.MAX_BYTES + 1))


def test_files_are_copied_and_ids_resolve_only_inside_the_folder(tmp_path):
    src = tmp_path / "notes.txt"
    src.write_text("hello")
    meta = attachments.add_file(str(src))
    assert Path(meta["path"]).parent == attachments.folder() and meta["name"] == "notes.txt"
    src.unlink()  # the copy survives the original being moved or deleted
    assert attachments.get(meta["id"])["name"] == "notes.txt"
    assert attachments.get("../../../etc/passwd") is None


def test_documents_become_text(tmp_path):
    d = docx.Document()
    d.add_paragraph("Weekly plan")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Mon", "Chest"
    d.save(tmp_path / "plan.docx")
    text = attachments.to_text(attachments.add_file(str(tmp_path / "plan.docx")))
    assert text.startswith("[Attached document: plan.docx]") and "Weekly plan" in text and "Mon | Chest" in text

    (tmp_path / "long.md").write_text("x" * 20000)
    assert "truncated" in attachments.to_text(attachments.add_file(str(tmp_path / "long.md")))


def test_images_go_through_the_vision_model(tmp_path):
    from PIL import Image
    Image.new("RGB", (3000, 2000), "white").save(tmp_path / "gym.png")
    vision = FakeVision()
    text = attachments.to_text(attachments.add_file(str(tmp_path / "gym.png")), vision_llm=vision)
    assert text.startswith("[Attached image: gym.png]\nGYM SCHEDULE")
    sent = vision.seen[0][1]["image_url"]["url"]
    assert sent.startswith("data:image/jpeg;base64,")  # downscaled and re-encoded before sending


def test_unreadable_attachment_is_reported_not_fatal(tmp_path):
    bad = attachments.add_bytes("broken.pdf", b"not really a pdf")
    assert "couldn't be read" in attachments.to_text(bad)


def test_history_and_titles_show_what_the_user_typed(aura_agent):
    meta = attachments.add_bytes("gym.png", b"\x89PNG")
    kwargs = {"display_text": "", "attachments": [{k: meta[k] for k in ("id", "name", "kind", "path")}]}
    aura_agent.graph.update_state(aura_agent.config, {"messages": [
        HumanMessage("[Attached image: gym.png]\nMon: Chest 7 PM", additional_kwargs=kwargs),
        AIMessage("Got your schedule!")]}, as_node="model")
    row = aura_agent.history()[0]
    assert row[0] == "user" and row[1] == "" and row[2][0]["name"] == "gym.png"  # not the transcription
    threads.touch(aura_agent.graph, aura_agent.thread_id)
    assert threads.list_threads(aura_agent.graph)[0]["title"] == "📎 gym.png"
    assert threads.search(aura_agent.graph, "chest")  # but the transcription is searchable
